"""Typed, lossless ingestion of the two supplied per-game projection sheets."""

from __future__ import annotations

from dataclasses import asdict, dataclass, field
from hashlib import sha256
from math import isfinite
from pathlib import Path
import re
from typing import Any

from openpyxl import load_workbook


SHEETS = {
    "LeagueWinner Stats": ("Rank", "Player", "Pos", "Team", "GP", "MPG", "FG%", "FT%", "FGM", "FGA", "FTM", "FTA", "3PM", "PTS", "REB", "AST", "STL", "BLK", "TO", "Value"),
    "Hashtag Stats": ("R#", "PLAYER", "POS", "TEAM", "GP", "MPG", "FG%", "FT%", "3pm", "PTS", "TREB", "AST", "STL", "BLK", "TO", "TOTAL"),
}
SHOOTING = re.compile(r"^\s*(0(?:\.\d+)?|1(?:\.0+)?)\s*\(\s*(\d+(?:\.\d+)?)\s*/\s*(\d+(?:\.\d+)?)\s*\)\s*$")


@dataclass
class Issue:
    sheet: str
    row: int
    field: str
    code: str
    severity: str
    detail: str = ""


@dataclass
class Projection:
    source: str
    sheet: str
    row: int
    raw: dict[str, Any]
    rank: int | None = None
    name: str | None = None
    position: str | None = None
    team: str | None = None
    gp: float | None = None
    mpg: float | None = None
    fg_pct_reported: float | None = None
    ft_pct_reported: float | None = None
    fgm: float | None = None
    fga: float | None = None
    ftm: float | None = None
    fta: float | None = None
    three_ptm: float | None = None
    pts: float | None = None
    reb: float | None = None
    ast: float | None = None
    stl: float | None = None
    blk: float | None = None
    turnovers: float | None = None
    source_score: float | None = None
    issues: list[Issue] = field(default_factory=list)

    def as_dict(self) -> dict[str, Any]:
        return asdict(self)


@dataclass
class ImportBatch:
    file_sha256: str
    records: list[Projection]
    issues: list[Issue]
    row_counts: dict[str, int]


def _issue(record: Projection, field: str, code: str, severity: str = "error", detail: str = "") -> None:
    record.issues.append(Issue(record.sheet, record.row, field, code, severity, detail))


def _number(record: Projection, column: str, *, nonnegative: bool = True) -> float | None:
    value = record.raw.get(column)
    if value is None or value == "":
        _issue(record, column, "missing_value")
        return None
    if isinstance(value, bool) or not isinstance(value, (int, float)) or not isfinite(value):
        _issue(record, column, "invalid_number", detail=repr(value))
        return None
    number = float(value)
    if nonnegative and number < 0:
        _issue(record, column, "negative_value", detail=str(number))
        return None
    return number


def _text(record: Projection, column: str) -> str | None:
    value = record.raw.get(column)
    if not isinstance(value, str) or not value.strip():
        _issue(record, column, "missing_value", "warning" if column in {"Pos", "POS", "Team", "TEAM"} else "error")
        return None
    return value.strip()


def _embedded_shooting(record: Projection, column: str) -> tuple[float | None, float | None, float | None]:
    value = record.raw.get(column)
    match = SHOOTING.fullmatch(value) if isinstance(value, str) else None
    if match is None:
        _issue(record, column, "invalid_shooting_format", detail=repr(value))
        return None, None, None
    return tuple(map(float, match.groups()))


def _validate_shooting(record: Projection, prefix: str) -> None:
    made = getattr(record, prefix + "m")
    attempted = getattr(record, prefix + "a")
    reported = getattr(record, prefix + "_pct_reported")
    if made is None or attempted is None or reported is None:
        return
    if made < 0 or attempted < 0 or made > attempted:
        _issue(record, prefix.upper(), "invalid_shooting_volume", detail=f"{made}/{attempted}")
        return
    if not 0 <= reported <= 1:
        _issue(record, prefix.upper() + "%", "invalid_percentage", detail=str(reported))
        return
    if attempted == 0:
        if made != 0:
            _issue(record, prefix.upper(), "invalid_shooting_volume", detail=f"{made}/0")
        else:
            _issue(record, prefix.upper(), "zero_attempts", "warning")
            if reported != 0:
                _issue(record, prefix.upper() + "%", "rate_without_attempts", "warning", f"reported={reported}")
        return
    # Source volumes are rounded to tenths; the displayed rate is rounded to
    # 0.1 percentage point (LeagueWinner) or 0.001 (Hashtag).
    lower = max(0.0, made - 0.05) / (attempted + 0.05)
    upper = min(attempted + 0.05, made + 0.05) / max(attempted - 0.05, 1e-12)
    precision = 0.0005
    if reported + precision < lower - 1e-12 or reported - precision > upper + 1e-12:
        _issue(record, prefix.upper() + "%", "rate_volume_discrepancy", "warning", f"reported={reported}; ratio={made / attempted:.6f}")


def parse_row(sheet: str, row: int, raw: dict[str, Any]) -> Projection:
    """Return a record even when cells are invalid, with row-scoped issues."""
    if sheet not in SHEETS:
        raise ValueError(f"Unsupported sheet: {sheet}")
    lw = sheet == "LeagueWinner Stats"
    record = Projection("LeagueWinner" if lw else "Hashtag", sheet, row, raw)
    record.name = _text(record, "Player" if lw else "PLAYER")
    record.position = _text(record, "Pos" if lw else "POS")
    record.team = _text(record, "Team" if lw else "TEAM")
    rank = _number(record, "Rank" if lw else "R#")
    if rank is not None:
        if not rank.is_integer() or rank < 1:
            _issue(record, "Rank" if lw else "R#", "invalid_rank")
        else:
            record.rank = int(rank)
    record.gp = _number(record, "GP")
    record.mpg = _number(record, "MPG")
    if record.gp == 0:
        _issue(record, "GP", "zero_games", "warning")
    if record.gp is not None and record.gp > 82:
        _issue(record, "GP", "implausible_games", "warning")
    if record.mpg is not None and record.mpg > 48:
        _issue(record, "MPG", "implausible_minutes", "warning")
    if lw:
        fg = _number(record, "FG%")
        ft = _number(record, "FT%")
        record.fg_pct_reported = fg / 100 if fg is not None else None
        record.ft_pct_reported = ft / 100 if ft is not None else None
        for attr, column in (("fgm", "FGM"), ("fga", "FGA"), ("ftm", "FTM"), ("fta", "FTA")):
            setattr(record, attr, _number(record, column))
    else:
        record.fg_pct_reported, record.fgm, record.fga = _embedded_shooting(record, "FG%")
        record.ft_pct_reported, record.ftm, record.fta = _embedded_shooting(record, "FT%")
    hashtag_threes_column = "3PM" if "3PM" in raw else "3pm"
    for attr, column in (("three_ptm", "3PM" if lw else hashtag_threes_column), ("pts", "PTS"), ("reb", "REB" if lw else "TREB"), ("ast", "AST"), ("stl", "STL"), ("blk", "BLK"), ("turnovers", "TO")):
        setattr(record, attr, _number(record, column))
    record.source_score = _number(record, "Value" if lw else "TOTAL", nonnegative=False)
    _validate_shooting(record, "fg")
    _validate_shooting(record, "ft")
    return record


def import_workbook(path: Path) -> ImportBatch:
    digest = sha256(path.read_bytes()).hexdigest()
    workbook = load_workbook(path, read_only=True, data_only=True)
    if set(workbook.sheetnames) != set(SHEETS):
        raise ValueError(f"Unexpected worksheets: {workbook.sheetnames}")
    records: list[Projection] = []
    counts: dict[str, int] = {}
    for sheet_name, expected in SHEETS.items():
        sheet = workbook[sheet_name]
        rows = sheet.iter_rows(values_only=True)
        header = tuple(next(rows))
        normalized_header = tuple("3pm" if sheet_name == "Hashtag Stats" and column == "3PM" else column for column in header)
        if normalized_header != expected:
            raise ValueError(f"Unexpected columns in {sheet_name}: {header}")
        count = 0
        for row_number, cells in enumerate(rows, start=2):
            if all(value is None for value in cells):
                continue
            record = parse_row(sheet_name, row_number, dict(zip(header, cells, strict=True)))
            records.append(record)
            count += 1
        counts[sheet_name] = count
    workbook.close()
    return ImportBatch(digest, records, [issue for record in records for issue in record.issues], counts)
